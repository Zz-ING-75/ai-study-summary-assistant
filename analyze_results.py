from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = PROJECT_ROOT / "data" / "output" / "evaluation_results.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "output" / "analysis_report.xlsx"

SCORE_COLUMNS = [
    "知识点覆盖率得分",
    "事实忠实度得分",
    "简洁度得分",
    "结构清晰度得分",
    "综合平均分",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="汇总V1/V2评测结果并生成Excel分析报告。")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="评测结果文件路径。")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="分析报告输出路径。")
    return parser.parse_args()


def build_report(input_path: Path, output_path: Path) -> None:
    if not input_path.exists():
        raise FileNotFoundError(f"未找到评测结果：{input_path}")

    source = pd.read_excel(input_path, sheet_name="评测结果")
    required = {"摘要ID", "资料ID", "Prompt版本", "模型", *SCORE_COLUMNS}
    missing = sorted(required.difference(source.columns))
    if missing:
        raise ValueError(f"评测结果缺少列：{', '.join(missing)}")

    detail_columns = [
        "摘要ID",
        "资料ID",
        "模型",
        "Prompt版本",
        "响应时间秒",
        "输入Token",
        "输出Token",
        *SCORE_COLUMNS,
    ]
    details = source[detail_columns].copy()
    # 同一资料可能同时保留旧版和新版摘要ID；统计时只使用最新评测结果。
    details = details.drop_duplicates(
        subset=["模型", "Prompt版本", "资料ID"], keep="last"
    ).copy()

    # Prompt实验只使用DeepSeek，并仅比较同时拥有V1和V2的配对资料。
    deepseek_details = details[
        details["模型"].astype(str).str.lower().str.contains("deepseek")
    ].copy()
    paired_ids = (
        deepseek_details.groupby("资料ID")["Prompt版本"]
        .nunique()
        .loc[lambda counts: counts >= 2]
        .index
    )
    prompt_source = deepseek_details[deepseek_details["资料ID"].isin(paired_ids)].copy()
    summary = (
        prompt_source.groupby("Prompt版本", as_index=False)
        .agg(
            样本数=("摘要ID", "count"),
            知识点覆盖率平均分=("知识点覆盖率得分", "mean"),
            事实忠实度平均分=("事实忠实度得分", "mean"),
            简洁度平均分=("简洁度得分", "mean"),
            结构清晰度平均分=("结构清晰度得分", "mean"),
            综合平均分=("综合平均分", "mean"),
            平均响应时间秒=("响应时间秒", "mean"),
            平均输入Token=("输入Token", "mean"),
            平均输出Token=("输出Token", "mean"),
        )
        .sort_values("Prompt版本")
    )

    numeric_columns = summary.select_dtypes(include="number").columns
    summary[numeric_columns] = summary[numeric_columns].round(4)

    model_source = details[details["Prompt版本"] == "V2"].copy()
    model_summary = (
        model_source.groupby("模型", as_index=False)
        .agg(
            样本数=("摘要ID", "count"),
            知识点覆盖率平均分=("知识点覆盖率得分", "mean"),
            事实忠实度平均分=("事实忠实度得分", "mean"),
            简洁度平均分=("简洁度得分", "mean"),
            结构清晰度平均分=("结构清晰度得分", "mean"),
            综合平均分=("综合平均分", "mean"),
            平均响应时间秒=("响应时间秒", "mean"),
            平均输出Token=("输出Token", "mean"),
        )
        .sort_values("综合平均分", ascending=False)
    )
    model_numeric = model_summary.select_dtypes(include="number").columns
    model_summary[model_numeric] = model_summary[model_numeric].round(4)

    conclusion_rows: list[dict[str, str]] = []
    if {"V1", "V2"}.issubset(set(summary["Prompt版本"])):
        indexed = summary.set_index("Prompt版本")
        comparisons = {
            "知识点覆盖率": "知识点覆盖率平均分",
            "事实忠实度": "事实忠实度平均分",
            "简洁度": "简洁度平均分",
            "结构清晰度": "结构清晰度平均分",
            "综合表现": "综合平均分",
        }
        for metric, column in comparisons.items():
            delta = float(indexed.loc["V2", column] - indexed.loc["V1", column])
            direction = "提升" if delta > 0 else "下降" if delta < 0 else "持平"
            conclusion_rows.append(
                {"比较项目": metric, "自动结论": f"V2相较V1{direction}{abs(delta):.4f}"}
            )

        time_delta = float(indexed.loc["V2", "平均响应时间秒"] - indexed.loc["V1", "平均响应时间秒"])
        speed_text = "更慢" if time_delta > 0 else "更快" if time_delta < 0 else "相同"
        conclusion_rows.append(
            {"比较项目": "响应时间", "自动结论": f"V2平均{speed_text}{abs(time_delta):.4f}秒"}
        )

        token_delta = float(indexed.loc["V2", "平均输出Token"] - indexed.loc["V1", "平均输出Token"])
        token_text = "增加" if token_delta > 0 else "减少" if token_delta < 0 else "不变"
        conclusion_rows.append(
            {"比较项目": "输出长度", "自动结论": f"V2平均输出Token{token_text}{abs(token_delta):.2f}"}
        )

    conclusions = pd.DataFrame(conclusion_rows)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="V1_V2汇总", index=False)
        model_summary.to_excel(writer, sheet_name="模型对比_V2", index=False)
        details.to_excel(writer, sheet_name="评分明细", index=False)
        conclusions.to_excel(writer, sheet_name="自动结论", index=False)

    workbook = load_workbook(output_path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in sheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 32)

    summary_sheet = workbook["V1_V2汇总"]
    for row in summary_sheet.iter_rows(min_row=2, min_col=3, max_col=10):
        for cell in row:
            cell.number_format = "0.0000"

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Prompt V1/V2质量得分对比"
    chart.y_axis.title = "平均得分（0-1）"
    chart.x_axis.title = "Prompt版本"
    chart.height = 8
    chart.width = 16
    data = Reference(summary_sheet, min_col=3, max_col=7, min_row=1, max_row=summary_sheet.max_row)
    categories = Reference(summary_sheet, min_col=1, min_row=2, max_row=summary_sheet.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    summary_sheet.add_chart(chart, "A6")

    workbook.save(output_path)
    print(f"分析报告已生成：{output_path}")


def main() -> None:
    args = parse_args()
    build_report(args.input, args.output)


if __name__ == "__main__":
    main()
